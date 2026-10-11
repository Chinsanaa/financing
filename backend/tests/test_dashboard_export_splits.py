"""GET /dashboard/export: split transactions render a "Split (n)" category
column instead of a category name, in the xlsx build."""
from io import BytesIO
from openpyxl import load_workbook

USER_A = "user-a-id"


def _headers(make_token, sub=USER_A):
    return {"Authorization": f"Bearer {make_token(sub=sub)}"}


def test_export_renders_split_count_for_split_transactions(client, patch_jwks, make_token, fake_db):
    fake_db.seed("categories", [
        {"id": "cat-1", "user_id": USER_A, "name": "Groceries"},
        {"id": "cat-2", "user_id": USER_A, "name": "Household"},
    ])
    fake_db.seed("transactions", [
        {"id": "t1", "user_id": USER_A, "merchant": "Costco", "description": "",
         "amount": 120.0, "timestamp": "2026-06-01T00:00:00", "category_id": None,
         "categories": None, "is_split": True, "label_source": "override"},
        {"id": "t2", "user_id": USER_A, "merchant": "Cafe", "description": "",
         "amount": 15.0, "timestamp": "2026-06-01T00:00:00", "category_id": "cat-1",
         "categories": {"name": "Groceries"}, "is_split": False, "label_source": "rule"},
    ])
    fake_db.seed("transaction_splits", [
        {"id": "s1", "user_id": USER_A, "transaction_id": "t1", "category_id": "cat-1", "amount": 80.0},
        {"id": "s2", "user_id": USER_A, "transaction_id": "t1", "category_id": "cat-2", "amount": 40.0},
    ])

    response = client.get("/dashboard/export", headers=_headers(make_token))

    assert response.status_code == 200
    wb = load_workbook(BytesIO(response.content))
    ws = wb.active
    rows_by_merchant = {row[1]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows_by_merchant["Costco"][3] == "Split (2)"
    assert rows_by_merchant["Cafe"][3] == "Groceries"


def _row(tid, merchant, amount, day, **extra):
    row = {"id": tid, "user_id": USER_A, "merchant": merchant, "description": "", "amount": amount,
           "timestamp": f"2026-{day}T00:00:00", "category_id": None, "categories": None,
           "is_split": False, "label_source": "none"}
    row.update(extra)
    return row


def test_export_has_v5_columns_and_all_kinds(client, patch_jwks, make_token, fake_db):
    fake_db.seed("transactions", [
        _row("e1", "Cafe", 15.0, "06-02", category_id="cat-1", categories={"name": "Groceries"}, label_source="rule"),
        _row("r1", "Shop", -5.0, "06-03", category_id="cat-1", categories={"name": "Groceries"}, flow_subtype="refund"),
        _row("i1", "Employer", 1000.0, "06-04", kind="income", flow_subtype="salary"),
        _row("x1", "Friend", -20.0, "06-05", kind="transfer", flow_subtype="red_packet"),
    ])

    response = client.get("/dashboard/export", headers=_headers(make_token))

    ws = load_workbook(BytesIO(response.content)).active
    assert [c.value for c in ws[1]] == [
        "Date", "Merchant", "Description", "Category", "Kind", "Amount",
        "Home amount", "Home currency", "Label Source"]
    rows = {row[1]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["Cafe"][3:6] == ("Groceries", "Expense", 15.0)
    assert rows["Shop"][4] == "Expense (Refund)"
    # Non-expense Category is "" (openpyxl reads an empty string back as None).
    assert rows["Employer"][3:6] == (None, "Income (Salary)", 1000.0)
    assert rows["Friend"][4] == "Transfer (Red packet)"
    assert rows["Cafe"][6] is None and rows["Cafe"][7] is None  # no home currency
    assert rows["Cafe"][8] == "rule"
    data_rows = list(ws.iter_rows(min_row=2))
    assert all(r[5].number_format == "#,##0.00" and r[6].number_format == "#,##0.00" for r in data_rows)


def test_export_home_amounts(client, patch_jwks, make_token, fake_db, fake_fx):
    from datetime import date
    fake_fx.home[USER_A] = "USD"
    fake_fx.rates[date(2026, 6, 1)] = 0.2
    fake_db.seed("transactions", [
        _row("e1", "Cafe", 15.0, "06-02"),
        _row("e2", "Old", 9.0, "01-02"),
    ])

    ws = load_workbook(BytesIO(client.get("/dashboard/export", headers=_headers(make_token)).content)).active

    rows = {row[1]: row for row in ws.iter_rows(min_row=2, values_only=True)}
    assert rows["Cafe"][6:8] == (3.0, "USD")
    assert rows["Old"][6:8] == (None, "USD")  # no strict rate → empty cell
